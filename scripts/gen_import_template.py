#!/usr/bin/env python3
"""Builds the product import template (Products → Import from Excel):
apps/web/public/templates/pharmatrack-product-import-template.xlsx

The header names are what apps/web/lib/inventory/importRows.ts auto-maps —
change them together. Categories come from packages/core/src/catalog/taxonomyRules.ts.
Also importable (`write_workbook`) to produce a pre-filled sheet for a pharmacy.

Usage: python3 scripts/gen_import_template.py   (needs: pip install openpyxl)
"""
from __future__ import annotations
import os, re, sys
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..")
OUT = os.path.join(ROOT, "apps/web/public/templates/pharmatrack-product-import-template.xlsx")

def load_taxonomy() -> dict[str, list[str]]:
    src = open(os.path.join(ROOT, "packages/core/src/catalog/taxonomyRules.ts"), encoding="utf8").read()
    block = src[src.index("export const DEFAULT_TAXONOMY"):src.index("} as const")]
    out = {}
    for m in re.finditer(r'^\s*"([^"]+)":\s*\[(.*)\],\s*$', block, re.M):
        out[m.group(1)] = re.findall(r'"([^"]+)"', m.group(2))
    return out

FORMS = ["Tablet", "Capsule", "Syrup", "Suspension", "Solution", "Drops", "Eye Drops", "Ear Drops", "Nasal Spray", "Cream", "Ointment",
         "Gel", "Lotion", "Injection", "Inhaler", "Sachet", "Powder", "Pessary", "Suppository", "Lozenge", "Spray", "Soap", "Device", "Other"]
UNITS = ["tablet", "capsule", "bottle", "tube", "vial", "ampoule", "sachet", "piece", "pessary", "suppository", "inhaler", "pack", "ml", "g"]
PACKS = ["Box", "Strip", "Blister", "Bottle", "Tube", "Vial", "Ampoule", "Sachet", "Piece", "Tin", "Packet", "Pack"]

# (header, width, required, help)
COLUMNS = [
    ("Product name", 42, True, "As printed on the pack: brand + strength + form, e.g. 'Augmentin 625mg Tablets'. One row per strength / form."),
    ("Generic name (INN)", 26, False, "e.g. Amoxicillin + Clavulanic acid. Leave blank and PharmaTrack fills it from the PPB register when it recognises the product."),
    ("Strength", 12, False, "e.g. 500mg, 125mg/5ml, 1%"),
    ("Dosage form", 13, False, "Pick from the list, or type your own."),
    ("Barcode (GTIN)", 16, False, "Scan the box barcode into this cell (8/12/13/14 digits). The column is formatted as text so Excel keeps every digit."),
    ("Old item code", 13, False, "Your previous system's code (e.g. A1226). Kept on the product so staff can still search / scan it."),
    ("Category", 24, False, "Pick from the list (PharmaTrack's default categories) or type your own. Blank → suggested for you."),
    ("Subcategory", 26, False, "Pick from the list or type your own."),
    ("Smallest unit sold", 14, True, "What you sell ONE of: tablet, capsule, bottle, tube, piece…"),
    ("Pack name", 11, False, "Box, Strip, Bottle… (how the supplier sells it to you)."),
    ("Units per pack", 10, True, "How many 'smallest units' are in ONE pack. A syrup or cream sold whole = 1."),
    ("Cost price per PACK (KES)", 14, True, "What you pay the supplier for ONE pack — straight from the invoice. Numbers only."),
    ("Selling price per UNIT (KES)", 14, True, "The price of ONE smallest unit (one tablet, one bottle). Numbers only."),
    ("Selling price per PACK (KES)", 14, False, "Only if a full pack has its own price (e.g. a box of 14 for 1,650 instead of 14 × 130). Becomes a pack size at the till."),
    ("Qty in stock (UNITS)", 11, False, "Counted in SMALLEST units: 3 boxes × 30 tablets = 90. Needs an expiry date to be loaded as stock."),
    ("Batch number", 13, False, "From the pack. Blank → 'OPENING'."),
    ("Expiry date", 12, False, "Needed for stock. MM/YYYY is fine (e.g. 03/2028) — read as the end of that month."),
    ("Reorder level (UNITS)", 11, False, "Warn me when stock falls to this many units. Blank → 10."),
    ("Prescription only? (Y/N)", 11, False, "Y for POM. Blank → taken from the PPB register when recognised."),
    ("Controlled? (Y/N)", 11, False, "Y for PPB controlled / psychotropic substances (go to the controlled-substances register). Blank → from the register."),
    ("Supplier", 18, False, "Your usual supplier for it. New names are added to your suppliers."),
]
NOTE_COL = ("Review note", 60, False, "Ignored by the import — for your own notes.")
REQ = PatternFill("solid", fgColor="B45309"); OPT = PatternFill("solid", fgColor="0F766E"); NOTE = PatternFill("solid", fgColor="6B7280")
EXAMPLES = [
    ["Augmentin 625mg Tablets", "Amoxicillin + Clavulanic acid", "625mg", "Tablet", "5000123456789", "A1001", "Antibiotics & Infections", "Penicillins",
     "tablet", "Box", 14, 1450, 130, 1650, 42, "B23K07", "06/2027", 14, "Y", "N", "Harleys Ltd"],
    ["Panadol Extra Tablets", "Paracetamol + Caffeine", "500mg/65mg", "Tablet", "", "", "Pain & Fever", "Paracetamol & combinations",
     "tablet", "Box", 100, 420, 10, "", 300, "PX2291", "11/2027", 50, "N", "N", ""],
    ["Ascoril Expectorant 100ml", "", "", "Syrup", "", "", "Cough, Cold & Allergy", "Cough syrups", "bottle", "Bottle", 1, 215, 350, "", 6, "AS901", "01/2028", 3, "", "", ""],
]

def write_workbook(path: str, rows: list[list] | None = None, notes: list[str] | None = None,
                   needs_fixing: list[list] | None = None, needs_header: list[str] | None = None, title: str = "PharmaTrack product import") -> None:
    tax = load_taxonomy()
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Products"
    cols = COLUMNS + ([NOTE_COL] if notes is not None else [])
    for i, (h, w, req, help_) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=h + (" *" if req else ""))
        c.font = Font(bold=True, color="FFFFFF"); c.alignment = Alignment(wrap_text=True, vertical="center")
        c.fill = NOTE if h == "Review note" else (REQ if req else OPT)
        c.comment = Comment(help_, "PharmaTrack"); c.comment.width = 300; c.comment.height = 120
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 48; ws.freeze_panes = "B2"
    for r_i, r in enumerate(rows or [], start=2):
        for c_i, v in enumerate(r, start=1):
            ws.cell(row=r_i, column=c_i, value=v)
        if notes is not None and notes[r_i - 2]:
            ws.cell(row=r_i, column=len(COLUMNS) + 1, value=notes[r_i - 2])
            for c_i in range(1, len(cols) + 1): ws.cell(row=r_i, column=c_i).fill = PatternFill("solid", fgColor="FEF3C7")
    last = max(2000, len(rows or []) + 1)

    lists = wb.create_sheet("Lists")
    subs = sorted({s for v in tax.values() for s in v})
    for j, (name, vals) in enumerate([("Category", list(tax)), ("Subcategory", subs), ("Dosage form", FORMS), ("Unit", UNITS), ("Pack name", PACKS), ("Y/N", ["Y", "N"])], start=1):
        lists.cell(row=1, column=j, value=name).font = Font(bold=True)
        for i, v in enumerate(vals, start=2): lists.cell(row=i, column=j, value=v)
        lists.column_dimensions[get_column_letter(j)].width = 30
    ref = lambda col, n: f"Lists!${col}$2:${col}${n + 1}"
    header = [c[0] for c in COLUMNS]
    def dv(field, formula, strict=False):
        col = get_column_letter(header.index(field) + 1)
        v = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
        ws.add_data_validation(v); v.add(f"{col}2:{col}{last}")
    dv("Category", ref("A", len(tax))); dv("Subcategory", ref("B", len(subs))); dv("Dosage form", ref("C", len(FORMS)))
    dv("Smallest unit sold", ref("D", len(UNITS))); dv("Pack name", ref("E", len(PACKS)))
    dv("Prescription only? (Y/N)", ref("F", 2), True); dv("Controlled? (Y/N)", ref("F", 2), True)
    for field in ("Units per pack", "Cost price per PACK (KES)", "Selling price per UNIT (KES)", "Selling price per PACK (KES)", "Qty in stock (UNITS)", "Reorder level (UNITS)"):
        col = get_column_letter(header.index(field) + 1)
        v = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True,
                           errorTitle="Numbers only", error="Type a number — no 'KES' or commas.")
        ws.add_data_validation(v); v.add(f"{col}2:{col}{last}")
    gcol = get_column_letter(header.index("Barcode (GTIN)") + 1)
    for r in range(2, last + 1): ws[f"{gcol}{r}"].number_format = "@"

    how = wb.create_sheet("How to fill", 1)
    lines = [
        (title, "h1"), ("", None),
        ("Fill one row per product on the 'Products' sheet, then in PharmaTrack: Products → Import from Excel.", None),
        ("Orange headers (*) are required. Hover over any header for help. Up to 2,000 rows per file.", None), ("", None),
        ("THE THREE RULES THAT PREVENT MOST IMPORT MISTAKES", "h2"),
        ("1.  Cost price is per PACK — exactly as on the supplier invoice.", None),
        ("2.  Selling price is per smallest UNIT — one tablet, one bottle.", None),
        ("3.  Stock quantity is in smallest UNITS — 3 boxes × 30 tablets = 90.", None),
        ("PharmaTrack works out the unit cost (pack cost ÷ units per pack) and warns you when a cost isn't below its selling price.", None), ("", None),
        ("CHECKED AGAINST THE PPB REGISTER", "h2"),
        ("Every product is looked up in the Pharmacy and Poisons Board register and our catalogue. Blank generic name, dosage form,", None),
        ("prescription / controlled flags and category are filled in for products it recognises — anything you type always wins.", None), ("", None),
        ("GOOD TO KNOW", "h2"),
        ("•  Same product with two expiry dates? Two rows with the same name and different batch / expiry → one product, two batches.", None),
        ("•  No expiry date yet? Leave it blank — the product is added, and you receive its stock later from Inventory → Receive.", None),
        ("•  Numbers only in price and quantity columns: 1200, not 'KES 1,200'.", None),
        ("•  Re-importing the same file never creates duplicates — products you already have are recognised.", None), ("", None),
        ("EXAMPLES (copy the layout — don't import these rows)", "h2"),
    ]
    for text, style in lines:
        how.append([text])
        c = how.cell(row=how.max_row, column=1)
        c.font = Font(bold=True, size=15) if style == "h1" else Font(bold=True, color="0F766E") if style == "h2" else Font(size=11)
    how.append(header)
    for c in how[how.max_row]: c.font = Font(bold=True)
    for ex in EXAMPLES:
        how.append(ex)
        for c in how[how.max_row]: c.font = Font(italic=True, color="4B5563")
    how.column_dimensions["A"].width = 44
    for i in range(2, len(header) + 1): how.column_dimensions[get_column_letter(i)].width = 16

    if needs_fixing:
        nf = wb.create_sheet("Needs fixing", 1)
        nf.append(needs_header)
        for c in nf[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="B91C1C")
        for r in needs_fixing: nf.append(r)
        for i in range(1, len(needs_header) + 1): nf.column_dimensions[get_column_letter(i)].width = 18 if i > 1 else 44
        nf.column_dimensions[get_column_letter(len(needs_header))].width = 70
        nf.freeze_panes = "B2"

    wb.active = 0
    wb.save(path)

if __name__ == "__main__":
    write_workbook(OUT)
    print("wrote", os.path.relpath(OUT, ROOT))
